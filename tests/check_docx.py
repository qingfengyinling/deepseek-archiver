"""结构校验：用 python-docx 重新打开生成的 docx，检查关键内容是否完整。"""
import os
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

out = Path(os.environ.get("DS_DIR") or (Path(__file__).parent / "out"))


def check(path: Path) -> None:
    print(f"\n=== {path.name} ({path.stat().st_size/1024:.1f} KB) ===")
    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
        names = z.namelist()
        print("zip 成员:", len(names), "| 损坏:", bad)
        dup = len(names) - len(set(names))
        print("重复成员:", dup)
        for must in ["[Content_Types].xml", "word/document.xml", "word/styles.xml",
                     "_rels/.rels", "word/_rels/document.xml.rels"]:
            assert must in names, must

    doc = Document(path)
    paras = doc.paragraphs
    print("段落:", len(paras), "| 表格:", len(doc.tables))
    styles = {}
    for p in paras:
        styles[p.style.name] = styles.get(p.style.name, 0) + 1
    print("样式分布:", dict(sorted(styles.items(), key=lambda kv: -kv[1])))

    # 超链接关系
    rels = doc.part.rels
    links = [r for r in rels.values() if r.reltype.endswith("/hyperlink")]
    print("外部超链接:", len(links), [r.target_ref for r in links][:3])

    # 代码块字体检查（字体定义在样式里，run 上不重复声明）
    cb = None
    for cand in ("Code Block", "DS Code Block"):
        try:
            cb = doc.styles[cand]
            break
        except KeyError:
            continue
    rpr = cb.element.find(qn("w:rPr")) if cb is not None else None
    fonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
    print("代码块样式字体:", fonts.get(qn("w:ascii")) if fonts is not None else None)

    if doc.tables:
        t = doc.tables[0]
        print("首个表格:", len(t.rows), "行 ×", len(t.columns), "列 | 表头:",
              [c.text for c in t.rows[0].cells])

    text = "\n".join(p.text for p in paras)
    for probe in ["核心结论", "计费对比", "import openai", "待办", "官方文档"]:
        if probe in text:
            print(f"  含内容: {probe}")
    print("正文首 200 字:", text[:200].replace("\n", " ⏎ "))

    # 检查所有 XML 部件是否可解析
    from lxml import etree
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.endswith(".xml") or n.endswith(".rels"):
                etree.fromstring(z.read(n))
    print("XML 全部可解析 ✓")


for name in sys.argv[1:] or ["session1.docx", "session2.docx", "merged.docx"]:
    check(out / name)

print("\n全部校验通过")
